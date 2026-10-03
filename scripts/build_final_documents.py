"""Build final Word sources and traceable figures from audited CSVs.

Use the Codex bundled Python. PDF conversion is performed by the companion
PowerShell script using a private, hidden Word COM instance.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import argparse
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports/generated/final_figures"
FONT = FontProperties(fname="C:/Windows/Fonts/msyh.ttc")
VERSION = "final-v1.0"
SOURCES = {}


def rows(path):
    with (ROOT / path).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, data):
    with (ROOT / path).open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(data[0]))
        w.writeheader()
        w.writerows(data)


def chart_font(fig):
    for text in fig.findobj(matplotlib.text.Text):
        text.set_fontproperties(FONT)
    fig.tight_layout(pad=1.2)


def figures():
    OUT.mkdir(parents=True, exist_ok=True)
    metrics = rows("data/public/company_quarterly_metrics.csv")
    q = sorted({r["period_end"] for r in metrics})
    labels = [f'{r[:4]} Q{(int(r[5:7])-1)//3+1}' for r in q]
    by = {(r["period_end"], r["metric_id"]): r for r in metrics}
    vals = lambda m: [float(by[(p, m)]["value"]) for p in q]
    x = range(len(q))
    fig, axes = plt.subplots(2, 1, figsize=(8, 4.3), sharex=True)
    for m, lab, marker, color in [("M001", "DAU 日均", "o", "#245B84"), ("M002", "MAU 月均", "s", "#4F7872")]:
        axes[0].plot(x, vals(m), marker=marker, label=lab, color=color, lw=2)
    axes[0].set(ylim=(0, 160), ylabel="百万用户")
    axes[0].legend(frameon=False, loc="upper left", ncol=2)
    axes[1].plot(x, [v * 100 for v in vals("M008")], "D-", color="#8E6520", lw=2)
    axes[1].set(ylim=(0, 50), ylabel="DAU/MAU 代理 %", xticks=list(x), xticklabels=labels)
    axes[0].annotate("2024 Q3 Video Call在Max推出", xy=(0, 37.2), xytext=(.8, 61), fontsize=8,
                     arrowprops={"arrowstyle": "->", "color": "#555555"})
    for ax in axes:
        ax.grid(axis="y", alpha=.18)
        ax.spines[["top", "right"]].set_visible(False)
    chart_font(fig)
    fig.savefig(OUT / "company_engagement.png", dpi=200)
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(8, 3.55))
    for m, lab, marker, color in [("M005", "订阅 bookings", "s", "#245B84"), ("M009", "订阅确认收入", "o", "#8E6520")]:
        axes[0].plot(x, vals(m), marker=marker, color=color, label=lab, lw=2)
    axes[0].set(ylim=(0, 330), ylabel="季度 USD 百万", xticks=list(x), xticklabels=labels)
    axes[0].tick_params(axis="x", rotation=55, labelsize=8)
    axes[0].legend(frameon=False, fontsize=8)
    axes[1].plot(x, vals("M004"), "D-", color="#4F7872", lw=2)
    axes[1].set(ylim=(0, 15), ylabel="期末付费账户 百万", xticks=list(x), xticklabels=labels)
    axes[1].tick_params(axis="x", rotation=55, labelsize=8)
    for ax in axes:
        ax.grid(axis="y", alpha=.18)
        ax.spines[["top", "right"]].set_visible(False)
    chart_font(fig)
    fig.savefig(OUT / "company_commercial.png", dpi=200)
    plt.close(fig)

    bridge = rows("data/public/stage3_base_period_bridge.csv")
    fig, ax = plt.subplots(figsize=(8, 3.1))
    bx = list(range(3))
    current = [float(r["current_qoq_log_points"]) for r in bridge]
    base = [float(r["base_period_contribution_log_points"]) for r in bridge]
    ax.bar([i - .18 for i in bx], current, width=.36, label="本年 Q1 到 Q2", color="#245B84")
    ax.bar([i + .18 for i in bx], base, width=.36, label="减去上年同期走势", color="#8E6520")
    for i, (a, b) in enumerate(zip(current, base)):
        ax.text(i-.18, a+.12, f"{a:+.2f}", ha="center", fontsize=9)
        ax.text(i+.18, b + (.12 if b >= 0 else -.48), f"{b:+.2f}", ha="center", fontsize=9)
    ax.axhline(0, color="#333", lw=.6)
    ax.set(xticks=bx, xticklabels=["DAU", "MAU", "DAU/MAU 代理"], ylim=(-5, 5.6), ylabel="100 × 自然对数增长点")
    ax.legend(frameon=False, ncol=2, loc="upper right")
    ax.spines[["top", "right"]].set_visible(False)
    chart_font(fig)
    fig.savefig(OUT / "period_bridge.png", dpi=200)
    plt.close(fig)

    coverage = rows("data/public/stage4_learning_coverage.csv")
    fig, ax = plt.subplots(figsize=(8, 2.45))
    for i, r in enumerate(coverage):
        n, a = int(r["assigned_n"]), int(r["analysed_n"])
        ax.barh(i, a, color="#245B84", height=.48, label="纳入后测分析" if i == 0 else None)
        ax.barh(i, n-a, left=a, color="#CAA975", height=.48, label="未纳入分析" if i == 0 else None)
        ax.text(a / 2, i, str(a), color="white", va="center", ha="center", fontsize=12)
        ax.text(a+(n-a)/2, i, str(n-a), va="center", ha="center", fontsize=11)
    ax.set(yticks=[0, 1], yticklabels=["Video Call", "Control"], xlim=(0, 350), xlabel="原随机分配人数 每组329")
    ax.invert_yaxis()
    ax.legend(frameon=False, loc="lower left", bbox_to_anchor=(0, 1.02), ncol=2)
    ax.spines[["top", "right"]].set_visible(False)
    chart_font(fig)
    fig.savefig(OUT / "learning_analysis_coverage.png", dpi=200)
    plt.close(fig)

    lineage = []
    for name, mids in [("F01", ("M001", "M002", "M008")), ("F02", ("M004", "M005", "M009"))]:
        for p in q:
            for m in mids:
                r = by[(p, m)]
                lineage.append({"figure_id": name, "input_table": "data/public/company_quarterly_metrics.csv", "key": f"{p}|{m}", "raw_value": r["value"], "unit": r["unit"], "display_transform": "100*ratio" if m == "M008" else "identity", "source_id": r["source_id"], "source_url": r["source_url"], "boundary": "global company; no feature causal attribution"})
    for r in bridge:
        for f in ("current_qoq_log_points", "base_period_contribution_log_points"):
            lineage.append({"figure_id": "F03", "input_table": "data/public/stage3_base_period_bridge.csv", "key": f'{r["metric_id"]}|{f}', "raw_value": r[f], "unit": "100 natural-log growth points", "display_transform": "2 decimal presentation only", "source_id": r["source_ids"], "source_url": r["source_urls"], "boundary": "arithmetic identity; not seasonal or causal adjustment"})
    for r in coverage:
        for f in ("analysed_n", "not_in_analysis_n"):
            lineage.append({"figure_id": "F04", "input_table": "data/public/stage4_learning_coverage.csv", "key": f'{r["study_id"]}|{r["arm"]}|{f}', "raw_value": r[f], "unit": "people", "display_transform": "identity", "source_id": "S015", "source_url": r["source_url"], "boundary": "analysis coverage; not all attrition or effect estimate"})
    write_csv("reports/figure_lineage.csv", lineage)
    return by


def set_font(run, size=None, bold=None):
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def setup(title):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin, section.bottom_margin = Inches(.65), Inches(.65)
    section.left_margin, section.right_margin = Inches(.72), Inches(.72)
    section.header_distance, section.footer_distance = Inches(.25), Inches(.3)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        s = doc.styles[name]
        s.font.name = "Microsoft YaHei"
        s._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15
    # The bundled default template contains a blue Title border. Remove the
    # inherited border explicitly; the visual hierarchy uses black typography.
    for style in doc.styles:
        ppr = style._element.find(qn("w:pPr"))
        if ppr is not None:
            for border in list(ppr.findall(qn("w:pBdr"))):
                ppr.remove(border)
    for name, size in [("Title", 24), ("Heading 1", 18), ("Heading 2", 12), ("Heading 3", 11)]:
        s = doc.styles[name]
        s.font.size = Pt(size)
        s.paragraph_format.space_before = Pt(9)
        s.paragraph_format.space_after = Pt(8)
    f = section.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(f.add_run(f"Duolingo Speaking 分析  {VERSION}  |  "), 8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    f._p.append(fld)
    doc.core_properties.title = title
    doc.core_properties.author = "Quanyi Liu"
    doc.core_properties.subject = "独立公开资料商业分析及条件性验证方案"
    return doc


def para(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    # Preserve text but remove Markdown emphasis markers for business prose.
    for i, part in enumerate(re.split(r"(\*\*.*?\*\*)", text)):
        bold = part.startswith("**") and part.endswith("**")
        set_font(p.add_run(part[2:-2] if bold else part.replace("`", "")), bold=bold)
    return p


def table(doc, headers, data, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    if widths is None:
        widths = [6.83 / len(headers)] * len(headers)
    for row_no, values in enumerate([headers] + data):
        row = t.rows[0] if row_no == 0 else t.add_row()
        # Rows expand and remain together; repeat header across natural breaks.
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))
        if row_no == 0:
            props.append(OxmlElement("w:tblHeader"))
        for i, val in enumerate(values):
            c = row.cells[i]
            c.width = Inches(widths[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "244B63" if row_no == 0 else ("F1F5F7" if row_no % 2 == 0 else "FFFFFF"))
            tcpr.append(shd)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{edge}")
                for k, v in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                    el.set(qn(f"w:{k}"), v)
                borders.append(el)
            tcpr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:w"), "95")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            tcpr.append(margins)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(val))
            set_font(r, 9, row_no == 0)
            if row_no == 0:
                r.font.color.rgb = RGBColor(255, 255, 255)
    para(doc, "")
    return t


def equation(doc, text):
    # Simple equations use native editable Word math; no raw LaTeX.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math = OxmlElement("m:oMath")
    mr = OxmlElement("m:r")
    mt = OxmlElement("m:t")
    mt.text = text
    mr.append(mt)
    math.append(mr)
    p._p.append(math)


def link(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement("w:r")
    rp = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "245B84")
    rp.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "17")
    rp.append(size)
    r.append(rp)
    txt = OxmlElement("w:t")
    txt.text = label
    r.append(txt)
    h.append(r)
    p._p.append(h)


def source(doc, text, ids=()):
    p = para(doc, text)
    for r in p.runs:
        set_font(r, 8.5)
        r.font.color.rgb = RGBColor.from_string("555555")
    for sid in ids:
        s = SOURCES[sid]
        link(doc, f'{sid}  {s["title"]}', s["url"])


def image(doc, name, caption):
    p = doc.add_paragraph()
    p.add_run().add_picture(str(OUT / name), width=Inches(6.75))
    source(doc, caption)


def page(doc, title):
    if len(doc.paragraphs) > 0:
        doc.add_page_break()
    doc.add_heading(title, level=1)


def business_book(by):
    doc = setup("Duolingo Speaking 扩张商业决策报告")
    doc.add_paragraph("Duolingo Speaking 扩张商业决策报告", style="Title")
    para(doc, "Executive Decision Book", style="Subtitle")
    para(doc, "作者 Quanyi Liu  |  独立公开资料分析  |  版本 final-v1.0\n资料截止 2026-10-02  |  财务窗口 2024 Q3 至 2026 Q2  |  制作 2026-10-03")
    doc.add_heading("先验证学习迁移 再判断权益扩张", level=1)
    para(doc, "**建议先审查目标账户现行权益与实施权限，优先验证H1的自然使用是否改善延后未练情境的口语表现。**当前证据足以形成可审查方案，尚不能批准扩大权益、确定盈利策略或预测用户增长。")
    para(doc, "H1指日语母语中级英语学习任务。现有研究支持特定条件下的短期学习信号，尚未验证自然使用、全员效果或经营回报。日语母语不等于现居日本，课程水平也不等于实际口语水平。")
    table(doc, ["发现", "商业含义", "下一动作"], [
        ["课程与AI对话为多家共同供给", "差异应在有效练习和迁移上验证", "固定用户任务和独立测评"],
        ["公司已披露向新Super扩展", "新增资格可能与现行权益重叠", "逐账户确认P0及真实增量"],
        ["贡献、成本与Max传播均未知", "无法判断扩张净值和总体排名", "按全员双臂收集账单与服务成本"],
    ], [2.15, 2.2, 2.48])
    para(doc, "**当前决策为准备验证，执行未就绪。**若目标Super已具备相同Lily权益，停止当前新增访问定义；不能为保留试验临时创造配额。外部自然任务研究可以检验另一种学习问题，但不能代替公司权益经营试验。")
    source(doc, "正文按市场、机制、经营、Driver、经济条件和执行方案展开。来源等级：A官方；C探索性体验/VOC；设计和会计恒等式分别标记。", ["S001", "S015"])

    page(doc, "市场证据支持研究问题 尚未支持产品胜负")
    para(doc, "**关键机会是帮助已有课程基础的用户独立表达，并把练习转移到未练任务。**Speak的日本英语课程与Free Talk、ELSA的Conversation Coach与Role-Play Lab说明相关供给已存在。Duolingo的优势应作为同任务待验证假设，不能由功能目录推定。")
    table(doc, ["评审维度", "已有证据", "决策所缺"], [
        ["任务与可达覆盖", "Lily、Speak、ELSA有英语对话供给", "同用户、设备、版本实际资格"],
        ["学习机制", "开放对话、课程内容、纠错提示", "延后未练任务的同口径效果"],
        ["体验与信任", "第一人称反例涉及起步、延迟和反馈", "可比故障率及纠错准确度"],
        ["激活与习惯", "课程节点和独立入口均可查", "曝光至有效开口及自然持续使用"],
        ["订阅与替代", "试用的付款和额度条件不同", "同SKU交易条件及Max替代净额"],
        ["成本与防御性", "课程加对话不是独占能力", "真实服务成本及可持续效果"],
    ], [1.25, 2.45, 3.13])
    para(doc, "**证据量须按单位理解。**竞品账本包含97条逐项观察、48个URL和34位可区分体验者；其中明确尝试目标功能的最多33位，Babbel明确亲测只有3位。它们不能合计为97名用户，亦无同口径满意度或学习效果排名。")
    para(doc, "用户确认来源的298条Duolingo自选态度只用于这一个产品。权益、质量、技术和价格主题帮助确定指标与反例；套餐未知200条、水平未知279条，不能可靠切成H1/H2或估计痛点发生率。")
    source(doc, "M01：market/competitor_evidence_v2.csv、competitor_scenario_review.md及两个quality_review；截止2026-09-29。计数为已审阅账本/提交标签，不是独立重编码或代表性抽样。")
    link(doc, "Speak 日本课程和目标语供给", "https://www.speak.com/jp/content")
    link(doc, "ELSA 新版 Coach 与 AI Chats 入口说明", "https://blog.elsaspeak.com/en/discover-the-new-elsa-speak-experience/")

    page(doc, "学习信号应转为全员迁移验证")
    para(doc, "**公开研究提供继续验证H1的理由，尚不足以给当前Super用户定标效果。**DRR-25-06研究对日语母语、B1.1英语课程成人分组，干预30天。通话组规定高频练习，对照组以课程练习为主；不同练习剂量和分析筛选限制自然权益政策外推。")
    image(doc, "learning_analysis_coverage.png", "F04 分析覆盖人数，源S015；329人/臂，263与304人纳入分析。未纳入比例20.1%与7.6%，并非全部失访。表stage4_learning_coverage.csv；构建build_final_documents.py。")
    table(doc, ["逻辑环节", "当前判定", "后续可检验问题"], [
        ["存在任务与可用功能", "提供方说明及探索线索", "该账户是否真的可进入"],
        ["通话与短期学习", "随机设计下完成者受限信号", "全员自然使用的延后迁移"],
        ["迁移与持续参与", "尚未识别", "不按使用依从或后测结果筛人"],
        ["参与与净商业回报", "尚未识别", "同窗贡献差、服务差与套餐传播"],
    ], [1.55, 2.25, 3.03])
    para(doc, "原PDF直接访问仍受限；同研究新闻稿和部分索引不等于全文复核，也不是独立重复研究。完成者成绩不能填留存、收入或转化uplift。Falstaff摘要主要为短语翻译信号，不能充当开放对话或H1的迁移证明。")
    source(doc, "研究为公司材料，证据A但外推受限。资料截至2026-10-02；居住地区、设备与当前套餐未确认。", ["S015", "S011"])

    page(doc, "公司活跃结构改善 功能贡献仍待识别")
    para(doc, "**窗口内DAU增长快于MAU，参与频次代理提高。**2024 Q3至2026 Q2，DAU由37.2增至58.7百万，MAU由113.1增至140.6百万，DAU/MAU由32.9%增至41.7%。这个跨季跨度不是同季同比，亦不是用户级留存。")
    image(doc, "company_engagement.png", "F01 公司全球指标，2024 Q3至2026 Q2，DAU日均/MAU月均，单位百万；代理为季度平均值比。源S004–S010/S002/S016；64行表中的相关指标；figure_lineage.csv逐点追溯。发布事件仅为背景，不表示因果。")
    para(doc, "公司披露CURR约84%，它描述近期活跃用户下一日留存，具体披露窗口不明，不能连成八季留存曲线，也不能用0.84的幂估计D30。")
    para(doc, "**决策影响：**公司经营背景支持同时观察学习与参与，但不能作为Video Call扩张成功证据。未来实验须保留原组、资格和曝光，区分新用户、持续用户及回流用户；June Streak Revival等整体活动仍是替代解释。")
    source(doc, "公开财务截至2026-06-30；产品及CURR陈述绑定2026-08-05披露日。A官方及分析师派生；不是功能因果样本。", ["S004", "S002", "S001"])

    page(doc, "购买金额与收入确认分化要求完整账单核算")
    image(doc, "company_commercial.png", "F02 全球公司季度USD百万及期末付费账户百万。bookings与收入为流量，账户为存量；2025 Q4订阅收入按年度减九个月流量差分。逐点源与转换见figure_lineage.csv、公司表input_lineage。")
    m = lambda p, k: float(by[(p, k)]["value"])
    yoy = lambda k: (m("2026-06-30", k) / m("2025-06-30", k) - 1) * 100
    table(doc, ["2026 Q2 指标", "水平", "显示值复算同比", "解释边界"], [
        ["期末付费账户", "12.7 百万", f'{yoy("M004"):.1f}%', "不能拆新增或续订"],
        ["订阅 bookings", f'{m("2026-06-30","M005"):.3f} USD百万', f'{yoy("M005"):.1f}%', "不是已确认收入"],
        ["订阅确认收入", f'{m("2026-06-30","M009"):.3f} USD百万', f'{yoy("M009"):.1f}%', "含历史购买确认"],
    ], [1.45, 1.9, 1.4, 2.08])
    para(doc, "**订阅权重使权益替代值得测量。**末季订阅收入占总收入约86.5%，但公司总额无法识别Super/Max价格、续订与用户结构。不能以期末账户为季度收入分母推造ARPU，亦不能把bookings与收入差额命名为利润。")
    para(doc, "**决策影响：**同一成熟计费窗内比较原分配账户的完整净贡献；升级、续订不重复记收入。现金贡献或会计贡献应在试验前选定，不能看结果后切口径。年度订阅尚无续订机会时不判长期续订收益。")
    source(doc, "A官方/派生同比；本期S002、原比较期S007。差分的2025 Q4原始两腿在公司表留存；舍入容差不是统计CI。", ["S002", "S007"])

    page(doc, "增长桥接区分当期变化与比较基数")
    para(doc, "**MAU同比项增强同时包含当期增长和上年比较基数。**2026 Q2相对Q1的MAU同比对数项增加3.48点，其中本年Q1至Q2变化2.01点，上年同期走势形成基数项1.47点。不能全部称为当期获客加速。")
    image(doc, "period_bridge.png", "F03 两组柱相加为同比对数项变化；DAU +1.49、MAU +3.48、代理 -1.99。单位100×自然对数增长点；不是普通百分点或因果贡献。stage3_base_period_bridge.csv，公式为本期环比对数项减去上年环比。")
    table(doc, ["诊断项", "证据与反证", "对验证的要求"], [
        ["活跃规模与频次", "算术方向稳健；基数、回流和构成可解释", "全员学习/参与分开测"],
        ["订阅购买与确认", "增速分化；价格、汇率和强基期为公司解释", "账单与套餐替代分开"],
        ["AI服务成本", "公司称成本压力及效率改善；无通话单位值", "失败/重试/购后成本全计"],
    ], [1.55, 2.6, 2.68])
    para(doc, "八季度不是64个独立时间样本。已有18个数值对象及2项事件独立来源复核，未完成全表双人源鉴证。舍入端点检查增强的是算术方向，不能提供渠道、用户构成或功能的因果识别。")
    source(doc, "D01：analysis/driver_evidence.csv与data_quality_report.md；F03源S002/S006/S007/S010/S016。管理层归因保留其身份。", ["S001"])

    page(doc, "经济评价保留未知并给出可检验条件")
    para(doc, "**当前无法估计扩张净值，能够定义正确的核算和盈亏平衡条件。**对同一人群、原套餐层、币种和成熟窗口，m为扣退款及必要非AI费用、未扣AI的账户贡献；c为完整服务成本。所有原分配账户进入均值，包含经确认的零使用者。")
    equation(doc, "Δv = Δm − Δc")
    equation(doc, "Δm ≥ Δc + F / N")
    para(doc, "第一式为每原分配账户的变量净差；第二式为单一匹配部署总体的盈亏平衡条件。Δm和Δc均为处理臂均值减对照臂均值；F只计新增固定投入，N为真实匹配部署人数，不能拿试验样本数或每千人归一化代替。")
    table(doc, ["敏感度关系", "商业用途", "不可外推"], [
        ["处理贡献 +1；对照贡献 -1", "两臂退款/费用/账单均需完整", "不是行为响应系数"],
        ["处理成本 -1；对照成本 +1", "抵扣已有成本，计失败和购后调用", "不等于成本实际下降"],
        ["固定费用增加减少总体净值", "只扣一次，不按虚构规模摊销", "未知F/N不给金额门槛"],
        ["可执行配额可约束赠送工作量", "真实账单单位映射后才定预算", "不能约束全部购后成本"],
    ], [2.05, 2.65, 2.13])
    para(doc, "完整政策还要把同一政策曝光下各原套餐层的变量差乘匹配部署人数，再减一次F。Super个体试验不能识别原Max传播；未测其他层不得设零。六组经济就绪表和两政策表的结果均为空，空值是未知。")
    para(doc, "**预测的合理收束：**保留条件关系，不给低/基准/高金额、DAU点预测、LTV、盈利概率或参数重要性排名。未来有真实均值和参数域才扩成敏感度范围。Teach Better项目须以同一政策曝光下全部增量服务、其他及固定投入对照真实批准预算，并满足学习护栏；未测层不设零，不把分数货币化，教学批准不等于盈利。")
    source(doc, "Q01：quant/evidence_model.py、evidence_inputs.json、evidence_assumptions.md与stage4_symbolic_sensitivity.csv。±1为会计恒等式常数；无经济模拟/实测效应。")

    page(doc, "三个权益方案的当前选择边界")
    para(doc, "**不选盈利赢家，先核实际参照和增量。**P0维持真实已有权益，P1为候选Super新增访问，P2为候选Free有限试用。H1/H2是用户任务，H3是净值假设，不能把三者列为互斥ROI策略。")
    table(doc, ["方案", "待验证价值", "主要风险", "进入或停止规则"], [
        ["P0 现行权益", "保持体验及套餐边界，提供共同参照", "现状可能损失学习机会", "按日期/课程/设备/账户核实；相对零非绝对成本零"],
        ["P1 Super增量", "自然口语迁移与Super保留价值", "重复权益、课程挤出、Max分流", "存在真实未开放子集且获权限才锁定；无增量停定义"],
        ["P2 Free试用", "接触有效练习与首次付费路径", "已有试用重复、无效调用和付费替代", "先核当前试用；未知成本/传播不批准全面扩张"],
    ], [.95, 1.6, 1.55, 2.73])
    para(doc, "公司8月披露多数新Super可用并计划向现有Super扩展，不证明10月目标iOS个人账户资格。旧Max不限次说明也不能覆盖新披露。以账户真实P0为对照，避免把已有覆盖重复计成新增市场。")
    para(doc, "日本App Store本轮可读列表有八个Super位置，金额包括¥9,900、¥11,400、¥14,200，但期限、个人/家庭和SKU未知。Max未出现于可读清单不证明不销售。当前不算套餐溢价、不年价除季度，也不把标价填净贡献。")
    source(doc, "O01：quant/strategy_comparison.csv；价格和资格逐项观察见stage4_public_observations.csv P01–P09/E01–E07。A官方供给和报价，非匹配合同。", ["S001", "S021"])

    page(doc, "一个优先学习问题及预设决策规则")
    para(doc, "**优先问题：真实新增Lily访问在自然使用下，是否提高H1用户延后未练情境的独立口语表现？**公司条件方案只在真实未覆盖Super子集、合法延后对照和实施权限成立后锁定；否则该定义停止在设计，P1尚未选定为试验。")
    table(doc, ["试验要素", "预设设计", "当前限制"], [
        ["处理与对照", "一个已核标准访问增量 vs 账户现行权益", "实际增量与权限未知"],
        ["随机化", "原合格账户1:1；资格/水平等前测分层", "设计建议，未分配真实用户"],
        ["唯一primary", "延后未练任务盲评总分，按原组ITT分析", "平行任务、评分与缺测策略待校准"],
        ["分析", "前测调整，效果量/区间/实际意义并报", "方差/MDE及业务受益门槛待核；不填N"],
        ["护栏", "课程挤出、失败、投诉、账单和服务成本", "界值/窗口/预算须试验前冻结"],
    ], [1.2, 3.2, 2.43])
    para(doc, "自然剂量试验估计新增访问的整体政策效果，不单独识别等剂量AI教学机制。未参加后测属于缺测，不打零分；未使用功能但完成后测仍保留原组。随机鼓励或志愿任务研究估计另一种效果，不能填公司权益收入模型。")
    para(doc, "**Scale：**只放行已验证的人群、处理及政策曝光范围，须满足学习实际意义、数据质量与护栏。教学预算须覆盖全部增量服务、其他及固定投入；盈利另需完整政策贡献/成本、Max传播及匹配N/F，未测层不设零。教学批准不等于盈利或全政策扩张批准。**Iterate：**机制可解释但效果未达或缺测仍影响判断，调整下一轮设计。**Stop：**无真实增量、无权限、严重伤害/成本越界或有证据显示无足够益处。未知不可读成成功。")
    source(doc, "E01：experiments/test_plan.md及event_spec.csv；完整N/MDE、估计对象、成熟窗口及评测规则见同版实验方案。这里为设计，无真实实验结果或CI。")

    page(doc, "缺口可转为数据请求和有限行动")
    table(doc, ["缺口类别", "所需最小证据", "影响与可继续工作"], [
        ["可公开但尚缺", "同条件账户资格、SKU期限、完整研究方法", "阻塞当前权益/价格比较；可继续任务设计"],
        ["实质内部或授权", "全员账单、退款、真实服务成本、原Max传播", "阻塞净值和政策Scale；能定义核算合同"],
        ["尚未产生的效果", "自然全员后测、延后迁移、续订成熟结果", "不能从更多评论补齐；需真实对照研究"],
        ["可合理推断", "课程对话共供给、学习机制/替代双向路径", "可选待证伪问题；不可确定方向或大小"],
    ], [1.5, 2.65, 2.68])
    para(doc, "**缺口严重影响真实商业回报板块，未使方法交接失效。**报告、核算、实验与监控设计已能彼此接合；实际执行反馈须在资格、权限与真实数据到位后产生。本项目未与Duolingo建立合作或发放权益。")
    para(doc, "监控以就绪与未来结果分开：历史公司趋势供背景；未知输入为未就绪，不能显示健康绿灯。未来使用预设Green/Yellow/Red条件和动作，真实阈值空白时阻止Scale。Excel可编辑图表及同版PDF仅为设计快照。")
    para(doc, "正式数据包只收必要数值事实、自己的汇总与计算、来源链接和方法，不含评论原文、个人数据、密钥、整份第三方PDF或退役D经济情景。CSV起可复现；网页人工核读不是自动刷新。标准Jupyter kernel仍未验证，已验证direct路径另列。")
    para(doc, "**后续接收者：**模拟产品经理先确认增量与权限；实验分析师冻结评测、样本和窗口；财务及工程完成账单/成本合同。缺一项就保留对应未知，不要求为填表补数字。")
    source(doc, "G01：quant/stage4_gap_register.csv 10项；完整交付入口docs/FINAL_DELIVERY.md。复盘仅覆盖实际研究/建模/制作，SOP规定输入、Gate与停止条件。")

    page(doc, "来源及复现索引")
    para(doc, "事实截至2026-10-02；动态权益与报价属于当次可读页面观察。来源总数、原子观察数和独立样本分开报告。下面索引和逐点lineage让复核者从图、表回到来源。")
    table(doc, ["对象", "权威数据或源文", "构建或检查"], [
        ["F01 活跃", "公司表 M001/M002/M008", "build_final_documents.py"],
        ["F02 商业", "公司表 M004/M005/M009", "build_final_documents.py"],
        ["F03 基期桥接", "stage3_base_period_bridge.csv", "rebuild_package_tables.py（复算）"],
        ["F04 分析覆盖", "stage4_learning_coverage.csv", "run_stage4_final.py"],
        ["M01/D01/Q01", "市场账本/driver/符号敏感度", "既有阶段质量记录及最终对账"],
        ["E01/G01", "实验方案/事件需求/缺口", "设计评审；不生成试验结果"],
    ], [1.3, 3.4, 2.13])
    for sid in ("S001", "S002", "S004", "S007", "S015", "S019", "S020", "S021"):
        s = SOURCES[sid]
        link(doc, f'{sid} {s["title"]}  |  {s["published_on"] or "日期未标"}', s["url"])
    source(doc, "逐图每个点的原值、单位、显示转换和URL：reports/figure_lineage.csv。市场正文见market_final_review.md，全部第三方链接及限制见data/source_register.csv与市场账本。")
    para(doc, "最终包包含版本、文件SHA-256、逐字段字典、共享判断与重建命令。文件校验只证明一致性，不认证网页、账户资格或产品因果。软件边界fixture不进入业务数据。")
    doc.save(ROOT / "reports/executive_decision_book.docx")


def markdown_to_docx(source_path, dest):
    text = (ROOT / source_path).read_text(encoding="utf-8")
    doc = setup("Duolingo Speaking 学习迁移试验方案")
    lines = text.splitlines()
    i = 0
    code = False
    def clean_cell(value):
        value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", value)
        return value.strip().replace("**", "").replace("`", "")
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            code = not code
            i += 1
            continue
        if line.startswith("|") and i+1 < len(lines) and re.match(r"^\|[ :\-|]+\|$", lines[i+1].strip()):
            headers = [clean_cell(v) for v in line.strip("|").split("|")]
            data = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                data.append([clean_cell(v) for v in lines[i].strip().strip("|").split("|")])
                i += 1
            if data:
                widths = [1.65, 5.18] if len(headers) == 2 else ([1.55, 2.55, 2.73] if len(headers) == 3 else None)
                table(doc, headers, data, widths)
            continue
        math_forms = {
            "`MDE_0": r"$\mathrm{MDE}_0(n)\approx (z_{1-\alpha/2}+z_{1-\beta})\sqrt{(\sigma^2_{T,\mathrm{res}}+\sigma^2_{C,\mathrm{res}})/n}$",
            "`n_obs": r"$n_{\mathrm{obs}}\approx\frac{(\sigma^2_{T,\mathrm{res}}+\sigma^2_{C,\mathrm{res}})(z_{1-\alpha/2}+z_{1-\beta})^2}{(\delta_A-\delta_L)^2}$",
            "`Δv_s": r"$\Delta v_s=(\overline{m}_{T,s}-\overline{m}_{C,s})-(\overline{c}_{T,s}-\overline{c}_{C,s})$",
        }
        math_form = next((v for k, v in math_forms.items() if line.startswith(k)), None)
        if math_form:
            fig = plt.figure(figsize=(8, .66))
            fig.text(.03, .45, math_form, fontsize=16, va="center")
            png = OUT / f"equation_{i}.png"
            fig.savefig(png, dpi=220, bbox_inches="tight", pad_inches=.06)
            plt.close(fig)
            p = doc.add_paragraph()
            p.add_run().add_picture(str(png), width=Inches(6.45))
        elif line.startswith("# "):
            para(doc, re.sub(r"[|：:｜]", " ", line[2:]), style="Title")
        elif line.startswith("## "):
            doc.add_heading(re.sub(r"[|：:｜]", " ", line[3:]), level=1)
        elif line.startswith("### "):
            doc.add_heading(re.sub(r"[|：:｜]", " ", line[4:]), level=2)
        elif line == "---":
            pass
        else:
            # Source Markdown retains formulas and links; Word has readable labels
            # and real clickable hyperlinks in addition to the source text.
            refs = re.findall(r"\[([^\]]+)\]\((https?://[^)]+)\)", line)
            clean = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", line)
            clean = re.sub(r"^[-*] ", "• ", clean)
            p = para(doc, clean)
            if code:
                for r in p.runs:
                    set_font(r, 9)
            for lab, url in refs:
                link(doc, lab, url)
        i += 1
    doc.save(ROOT / dest)


def main():
    global SOURCES
    SOURCES = {r["source_id"]: r for r in rows("data/source_register.csv")}
    ap = argparse.ArgumentParser()
    ap.add_argument("--test-plan-only", action="store_true")
    ap.add_argument("--book-only", action="store_true")
    args = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    if not args.test_plan_only:
        by = figures()
        business_book(by)
    if not args.book_only and (ROOT / "experiments/test_plan.md").exists():
        markdown_to_docx("experiments/test_plan.md", "experiments/test_plan.docx")
    print(json.dumps({"version": VERSION, "book_created": (ROOT / "reports/executive_decision_book.docx").exists(), "test_plan_created": (ROOT / "experiments/test_plan.docx").exists()}, ensure_ascii=False))


if __name__ == "__main__":
    main()
